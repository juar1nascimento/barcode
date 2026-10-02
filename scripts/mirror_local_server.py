"""Espelho local do Inventário GTI SESA.

Executado exclusivamente no servidor local. O PostgreSQL/Supabase é a origem
operacional; este agente materializa a cópia principal em disco local.

Destino padrão:
    \\172.17.27.246\\t.i\\02 - SUPORTE\\BACKUP\\Inventário gti-sesa

Segredos nunca ficam neste arquivo. Configure-os somente no ambiente do
serviço local.
"""

from __future__ import annotations

import hashlib
import json
import os
import tempfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from urllib.parse import quote

import openpyxl
import requests
from psycopg.rows import dict_row
import psycopg

COLUNAS_SITE = [
    "Setor",
    "Tipo de Patrimônio",
    "Nº de Patrimônio",
    "Fabricante",
    "Data Cadastro",
]

DEFAULT_ROOT = r"\\172.17.27.246\t.i\02 - SUPORTE\BACKUP\Inventário gti-sesa"
MIRROR_ROOT = Path(os.getenv("INVENTARIO_MIRROR_ROOT", DEFAULT_ROOT))
DB_CONNECT_TIMEOUT = int(os.getenv("INVENTARIO_MIRROR_DB_TIMEOUT", "15"))
HTTP_TIMEOUT = int(os.getenv("INVENTARIO_MIRROR_HTTP_TIMEOUT", "60"))
USER_AGENT = "Inventario-GTI-SESA-LocalMirror/1.0"


def _require(value: str | None, name: str) -> str:
    value = str(value or "").strip()
    if not value:
        raise RuntimeError(f"Configuração obrigatória ausente: {name}.")
    return value


def _db_connection() -> psycopg.Connection:
    database_url = os.getenv("DATABASE_URL")
    if database_url:
        return psycopg.connect(database_url, connect_timeout=DB_CONNECT_TIMEOUT, row_factory=dict_row)

    host = _require(os.getenv("PGHOST"), "PGHOST")
    port = os.getenv("PGPORT", "5432")
    database = _require(os.getenv("PGDATABASE"), "PGDATABASE")
    user = _require(os.getenv("PGUSER"), "PGUSER")
    password = _require(os.getenv("PGPASSWORD"), "PGPASSWORD")
    sslmode = os.getenv("PGSSLMODE", "require")
    return psycopg.connect(
        host=host,
        port=port,
        dbname=database,
        user=user,
        password=password,
        sslmode=sslmode,
        connect_timeout=DB_CONNECT_TIMEOUT,
        row_factory=dict_row,
    )


def _safe_sheet_name(name: str, used: set[str]) -> str:
    invalid = '[]:*?/\\'
    base = "".join("_" if c in invalid else c for c in str(name or "").strip()) or "Sem unidade"
    base = base[:31]
    candidate = base
    index = 2
    while candidate in used:
        suffix = f"_{index}"
        candidate = f"{base[:31-len(suffix)]}{suffix}"
        index += 1
    used.add(candidate)
    return candidate


def _atomic_write_bytes(destination: Path, payload: bytes) -> None:
    destination.parent.mkdir(parents=True, exist_ok=True)
    fd, tmp_name = tempfile.mkstemp(prefix=f".{destination.name}.", suffix=".tmp", dir=destination.parent)
    try:
        with os.fdopen(fd, "wb") as handle:
            handle.write(payload)
            handle.flush()
            os.fsync(handle.fileno())
        os.replace(tmp_name, destination)
    except Exception:
        try:
            os.unlink(tmp_name)
        except OSError:
            pass
        raise


def _atomic_write_json(destination: Path, data: Any) -> None:
    payload = json.dumps(data, ensure_ascii=False, indent=2, default=str).encode("utf-8")
    _atomic_write_bytes(destination, payload)


def _build_workbook(rows: list[dict[str, Any]]) -> bytes:
    workbook = openpyxl.Workbook()
    default = workbook.active
    workbook.remove(default)

    by_unit: dict[str, list[dict[str, Any]]] = {}
    for row in rows:
        by_unit.setdefault(str(row["unidade"]), []).append(row)

    used: set[str] = set()
    for unidade, itens in sorted(by_unit.items(), key=lambda item: item[0].casefold()):
        ws = workbook.create_sheet(_safe_sheet_name(unidade, used))
        ws.append(COLUNAS_SITE)
        for cell in ws[1]:
            cell.font = cell.font.copy(bold=True)
        for item in itens:
            ws.append([
                str(item["setor"] or ""),
                str(item["tipo"] or ""),
                str(item["numero_patrimonio"] or ""),
                str(item["fabricante"] or ""),
                str(item["data_cadastro"] or ""),
            ])
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        widths = [28, 24, 24, 30, 22]
        for idx, width in enumerate(widths, start=1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(idx)].width = width

    if not by_unit:
        ws = workbook.create_sheet("Inventario")
        ws.append(COLUNAS_SITE)

    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
        temp_path = Path(tmp.name)
    try:
        workbook.save(temp_path)
        return temp_path.read_bytes()
    finally:
        try:
            temp_path.unlink()
        except OSError:
            pass


def _download_storage_object(base_url: str, service_key: str, bucket: str, path: str) -> bytes:
    url = f"{base_url.rstrip('/')}/storage/v1/object/{quote(bucket, safe='')}/{quote(path, safe='/')}"
    response = requests.get(
        url,
        headers={"Authorization": f"Bearer {service_key}", "apikey": service_key, "User-Agent": USER_AGENT},
        timeout=HTTP_TIMEOUT,
    )
    response.raise_for_status()
    return response.content


def _mirror_photos(conn: psycopg.Connection, root: Path) -> tuple[int, int]:
    base_url = _require(os.getenv("SUPABASE_URL"), "SUPABASE_URL")
    service_key = _require(os.getenv("SUPABASE_SERVICE_ROLE_KEY"), "SUPABASE_SERVICE_ROLE_KEY")

    with conn.cursor() as cur:
        cur.execute(
            """select f.id, f.patrimonio_id, f.ordem, f.storage_bucket, f.storage_path,
                      f.arquivo_nome, f.sha256, p.numero_patrimonio
                 from public.patrimonio_fotos f
                 join public.patrimonios p on p.id=f.patrimonio_id
                order by p.numero_patrimonio, f.ordem, f.id"""
        )
        fotos = cur.fetchall()

    copied = 0
    for foto in fotos:
        payload = _download_storage_object(
            base_url,
            service_key,
            str(foto["storage_bucket"]),
            str(foto["storage_path"]),
        )
        digest = hashlib.sha256(payload).hexdigest()
        expected = str(foto["sha256"] or "").lower()
        if expected and digest != expected:
            raise RuntimeError(f"Integridade da foto {foto['id']} inválida: SHA-256 divergente.")

        patrimonio = str(foto["numero_patrimonio"])
        filename = Path(str(foto["arquivo_nome"] or f"foto_{foto['id']}.jpg")).name
        destination = root / "fotos" / "patrimonio" / patrimonio / f"{int(foto['ordem']):02d}_{filename}"
        _atomic_write_bytes(destination, payload)
        copied += 1

    return copied, len(fotos)


def run() -> dict[str, Any]:
    root = MIRROR_ROOT
    (root / "dados").mkdir(parents=True, exist_ok=True)
    (root / "fotos" / "patrimonio").mkdir(parents=True, exist_ok=True)
    (root / "manifest").mkdir(parents=True, exist_ok=True)
    (root / "backup").mkdir(parents=True, exist_ok=True)
    (root / "logs").mkdir(parents=True, exist_ok=True)

    started = datetime.now(timezone.utc)

    with _db_connection() as conn:
        with conn.cursor() as cur:
            cur.execute(
                """select p.id, p.numero_patrimonio, p.codigo_barras, p.tipo,
                          coalesce(p.fabricante,'') as fabricante,
                          p.data_cadastro, p.atualizado_em,
                          coalesce(s.nome,'') as setor,
                          u.nome as unidade,
                          p.ativo
                     from public.patrimonios p
                     join public.unidades u on u.id=p.unidade_id
                     left join public.setores s on s.id=p.setor_id
                    order by u.nome, s.nome, p.id"""
            )
            all_rows = cur.fetchall()

            cur.execute(
                """select id, nome, tipo, ativo, criado_em
                     from public.unidades order by id"""
            )
            unidades = cur.fetchall()

            cur.execute(
                """select id, unidade_id, nome, numero_consultorio, especialidade, ativo, criado_em
                     from public.setores order by unidade_id, id"""
            )
            setores = cur.fetchall()

            cur.execute(
                """select id, patrimonio_id, tipo, unidade_origem_id, setor_origem_id,
                          unidade_destino_id, setor_destino_id, motivo, observacao, usuario, criado_em
                     from public.movimentacoes_patrimonio order by id"""
            )
            movimentos = cur.fetchall()

            cur.execute(
                """select id, patrimonio_id, ordem, storage_bucket, storage_path,
                          arquivo_nome, mime_type, tamanho_bytes, largura, altura, sha256, criado_em
                     from public.patrimonio_fotos order by patrimonio_id, ordem, id"""
            )
            fotos = cur.fetchall()

    ativos = [row for row in all_rows if bool(row["ativo"])]
    workbook_payload = _build_workbook(ativos)
    workbook_path = root / "dados" / "inventario_site.xlsx"
    _atomic_write_bytes(workbook_path, workbook_payload)

    copied_photos, total_photos = _mirror_photos(
        _db_connection(),
        root,
    )

    archive = {
        "gerado_em": datetime.now(timezone.utc).isoformat(),
        "fonte_operacional": "Supabase PostgreSQL",
        "destino_principal": str(root),
        "patrimonios": all_rows,
        "unidades": unidades,
        "setores": setores,
        "movimentacoes": movimentos,
        "fotos": fotos,
    }
    _atomic_write_json(root / "dados" / "inventario_completo.json", archive)

    checksums = {
        "inventario_site.xlsx": hashlib.sha256(workbook_payload).hexdigest(),
    }
    for foto in (root / "fotos" / "patrimonio").rglob("*"):
        if foto.is_file():
            checksums[str(foto.relative_to(root)).replace("\\", "/")] = hashlib.sha256(foto.read_bytes()).hexdigest()
    _atomic_write_json(root / "manifest" / "checksums.json", checksums)

    finished = datetime.now(timezone.utc)
    manifest = {
        "status": "ok",
        "inicio": started.isoformat(),
        "fim": finished.isoformat(),
        "duracao_segundos": (finished - started).total_seconds(),
        "destino": str(root),
        "patrimonios_total": len(all_rows),
        "patrimonios_ativos": len(ativos),
        "fotos_total": total_photos,
        "fotos_mirroradas": copied_photos,
        "arquivo_tabela": str(workbook_path),
    }
    _atomic_write_json(root / "manifest" / "estado.json", manifest)
    return manifest


if __name__ == "__main__":
    resultado = run()
    print(json.dumps(resultado, ensure_ascii=False))
