import os
import re
from typing import Optional, Tuple, List, Any
import pandas as pd
import streamlit as st

# ==============================================================================
# CONFIGURAÇÕES E CONSTANTES DE DADOS
# ==============================================================================
ARQUIVO_EXCEL: str = "Tabela_Patrimonios_UBS_Feu_Rosa.xlsx"
GOOGLE_SHEET_URL: str = "https://docs.google.com/spreadsheets/d/12mNKTWLExRwZx3EKSB78oTScQk6ctGvi6eNKt5QyXEw/edit?usp=sharing"

COLUNA_CHAVE: str = "Local / Setor"

# Sufixo padrão exigido para colunas de patrimônio
SUFIXO_PATRIMONIO: str = " - N de Patrimônio"

COLUNAS_PADRAO: List[str] = [
    COLUNA_CHAVE,
    f"CPU{SUFIXO_PATRIMONIO}",
    f"Monitores{SUFIXO_PATRIMONIO}",
    f"Nobreak{SUFIXO_PATRIMONIO}",
    f"Teclado{SUFIXO_PATRIMONIO}",
    f"Mouse{SUFIXO_PATRIMONIO}",
    f"Impressora{SUFIXO_PATRIMONIO}"
]

COLUNAS_OBSOLETAS: List[str] = [
    "Patrimônio PC", "Patrimônio Tela", "Patrimônio Nobreak", 
    "➕ Outra descrição (Criar nova coluna ao final)", "cameras"
]

SETORES_PADRAO: List[str] = [
    "Consultório", "Gerência", "Administração", "Farmácia",
    "Almoxarifado", "Sala de Preparo", "Sala dos Agentes de Saúde",
    "Sala de Curativo", "Recepção", "Sala de Vacina"
]

LISTA_URS_PADRAO: List[str] = [
    "Selecione uma URS...",
    "URS Feu Rosa",
    "URS Jacaraípe",
    "URS Novo Horizonte",
    "URS Serra Dourada"
]

LISTA_UBS_PADRAO: List[str] = [
    "Selecione uma UBS...",
    "UBS Planalto Serrano",
    "UBS Bairro das Laranjeiras",
    "UBS Nova Carapina",
    "UBS Vila Nova de Colares",
    "UBS Porto Canoa"
]

# ==============================================================================
# FUNÇÕES AUXILIARES DE FORMATAÇÃO E SANITIZAÇÃO
# ==============================================================================
def formatar_nome_patrimonio(nome_patrimonio: str) -> str:
    """
    Garante a regra de nomenclatura inserindo ' - N de Patrimônio' no cabeçalho.
    Exemplo: 'CPU' -> 'CPU - N de Patrimônio'
    """
    nome_limpo = nome_patrimonio.strip()
    if not nome_limpo or nome_limpo == COLUNA_CHAVE or nome_limpo.startswith("Fabricante "):
        return nome_limpo
    if SUFIXO_PATRIMONIO.lower() not in nome_limpo.lower():
        return f"{nome_limpo}{SUFIXO_PATRIMONIO}"
    return nome_limpo

def sanitizar_nome_aba(nome_unidade: str) -> str:
    """Higieniza o nome da URS/UBS para ser um nome de aba válido."""
    if not nome_unidade or not str(nome_unidade).strip():
        return "Geral"
    nome_limpo = re.sub(r'[\\/*?:\[\]]', '_', str(nome_unidade).strip())
    return nome_limpo[:31]

# ==============================================================================
# CAMADA DE CONEXÃO REMOTA (GOOGLE SHEETS)
# ==============================================================================
def obter_conexao_gsheets() -> Optional[Any]:
    """Estabelece conexão com o Google Sheets usando a lib streamlit-gsheets."""
    try:
        from streamlit_gsheets import GSheetsConnection
        return st.connection("gsheets", type=GSheetsConnection)
    except Exception as err:
        print(f"[AVISO GSHEETS]: {err}")
        return None

def sincronizar_google_sheets(df: pd.DataFrame, nome_aba: str) -> bool:
    """Sincroniza os dados atualizados com o Google Sheets."""
    try:
        conn = obter_conexao_gsheets()
        if conn is not None:
            conn.update(spreadsheet=GOOGLE_SHEET_URL, worksheet=nome_aba, data=df)
            st.toast(f"☁️ Aba '{nome_aba}' sincronizada no Google Sheets!")
            return True
    except Exception as err:
        print(f"[ERRO AO SINCRONIZAR GSHEETS]: {err}")
        st.toast(f"⚠️ Salvo localmente. Erro ao sincronizar Google Sheets: {err}")
    return False

# ==============================================================================
# REGRAS DE NEGÓCIO E ORGANIZAÇÃO DAS COLUNAS DE FABRICANTE
# ==============================================================================
def padronizar_e_organizar_df(df: pd.DataFrame) -> pd.DataFrame:
    """
    Higieniza o DataFrame, aplica a regra de cabeçalho ' - N de Patrimônio'
    e agrupa cada fabricante ao lado de seu respectivo patrimônio.
    """
    colunas_invisiveis = [
        c for c in df.columns 
        if c in COLUNAS_OBSOLETAS or str(c).startswith("➕") or "Unnamed" in str(c)
    ]
    if colunas_invisiveis:
        df = df.drop(columns=colunas_invisiveis, errors="ignore")

    renomear_map = {}
    for col in df.columns:
        if col != COLUNA_CHAVE and not str(col).startswith("Fabricante "):
            col_formatada = formatar_nome_patrimonio(str(col))
            if col_formatada != col:
                renomear_map[col] = col_formatada

    if renomear_map:
        df = df.rename(columns=renomear_map)

    if COLUNA_CHAVE not in df.columns:
        df.insert(0, COLUNA_CHAVE, "")

    for col in COLUNAS_PADRAO:
        if col not in df.columns:
            df[col] = ""

    ordem_colunas = [COLUNA_CHAVE]
    todas_colunas = [str(c).strip() for c in df.columns]

    for col in todas_colunas:
        if col == COLUNA_CHAVE or col.startswith("Fabricante "):
            continue
        if col not in ordem_colunas:
            ordem_colunas.append(col)
        
        col_fab = f"Fabricante {col}"
        if col_fab in todas_colunas and col_fab not in ordem_colunas:
            ordem_colunas.append(col_fab)

    df_processado = df.reindex(columns=ordem_colunas).fillna("").astype(str)
    df_processado[COLUNA_CHAVE] = df_processado[COLUNA_CHAVE].str.strip()
    return df_processado

def aplicar_estilo_excel(caminho_arquivo: str, nome_aba: str) -> None:
    """Aplica formatação visual profissional à aba do arquivo Excel."""
    if not os.path.exists(caminho_arquivo):
        return
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter

        wb = openpyxl.load_workbook(caminho_arquivo)
        if nome_aba not in wb.sheetnames:
            return
        ws = wb[nome_aba]

        header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        row_even_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
        row_odd_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
        col_chave_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")

        thin_border_side = Side(border_style="thin", color="CBD5E1")
        border_box = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
        align_center = Alignment(horizontal="center", vertical="center")
        align_left = Alignment(horizontal="left", vertical="center")
        font_data = Font(name="Segoe UI", size=10, color="0F172A")
        font_sector = Font(name="Segoe UI", size=10, bold=True, color="0F172A")

        ws.row_dimensions[1].height = 28
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = align_center
            cell.border = border_box

        for r in range(2, ws.max_row + 1):
            ws.row_dimensions[r].height = 22
            current_fill = row_odd_fill if r % 2 == 1 else row_even_fill
            for c in range(1, ws.max_column + 1):
                cell = ws.cell(row=r, column=c)
                cell.border = border_box
                cell.number_format = "@"
                if c == 1:
                    cell.fill = col_chave_fill
                    cell.font = font_sector
                    cell.alignment = align_left
                else:
                    cell.fill = current_fill
                    cell.font = font_data
                    cell.alignment = align_center

        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 5, 18)

        wb.save(caminho_arquivo)
    except Exception as e:
        print(f"[AVISO FORMATAÇÃO EXCEL]: {e}")

# ==============================================================================
# CAMADA DE PERSISTÊNCIA E CRUD MULTI-ABAS
# ==============================================================================
@st.cache_data(ttl=60)
def carregar_dados_excel(unidade_nome: str = "Geral") -> Tuple[pd.DataFrame, List[str]]:
    """Carrega os dados exclusivamente da aba da URS ou UBS selecionada."""
    nome_aba = sanitizar_nome_aba(unidade_nome)
    
    if os.path.exists(ARQUIVO_EXCEL):
        try:
            xls = pd.ExcelFile(ARQUIVO_EXCEL)
            if nome_aba in xls.sheet_names:
                df = pd.read_excel(xls, sheet_name=nome_aba, dtype=str, keep_default_na=False)
            else:
                df = pd.DataFrame(columns=COLUNAS_PADRAO)
        except Exception:
            df = pd.DataFrame(columns=COLUNAS_PADRAO)
    else:
        df = pd.DataFrame(columns=COLUNAS_PADRAO)

    df = padronizar_e_organizar_df(df)
    
    houve_alteracao = False
    for col in COLUNAS_PADRAO:
        if col not in df.columns:
            df[col] = ""
            houve_alteracao = True

    for setor in SETORES_PADRAO:
        if not df[COLUNA_CHAVE].str.lower().eq(setor.lower()).any():
            nova_linha = {col: "" for col in df.columns}
            nova_linha[COLUNA_CHAVE] = setor
            df = pd.concat([df, pd.DataFrame([nova_linha])], ignore_index=True)
            houve_alteracao = True

    df = padronizar_e_organizar_df(df)
    if houve_alteracao or not os.path.exists(ARQUIVO_EXCEL):
        salvar_no_excel(df, unidade_nome=nome_aba)

    return df, list(df.columns)

def salvar_no_excel(df: pd.DataFrame, unidade_nome: str = "Geral") -> bool:
    """Persiste os dados no arquivo Excel local e aciona sincronização com o Google Sheets."""
    nome_aba = sanitizar_nome_aba(unidade_nome)
    df_limpo = padronizar_e_organizar_df(df)
    sucesso_local = False

    try:
        import openpyxl
        if not os.path.exists(ARQUIVO_EXCEL):
            with pd.ExcelWriter(ARQUIVO_EXCEL, engine="openpyxl") as writer:
                df_limpo.to_excel(writer, sheet_name=nome_aba, index=False)
        else:
            wb = openpyxl.load_workbook(ARQUIVO_EXCEL)
            if nome_aba in wb.sheetnames:
                del wb[nome_aba]
            wb.save(ARQUIVO_EXCEL)
            wb.close()

            with pd.ExcelWriter(ARQUIVO_EXCEL, engine="openpyxl", mode="a") as writer:
                df_limpo.to_excel(writer, sheet_name=nome_aba, index=False)

        aplicar_estilo_excel(ARQUIVO_EXCEL, nome_aba)
        carregar_dados_excel.clear()
        sucesso_local = True
    except Exception as err:
        st.error(f"❌ Erro ao salvar a aba '{nome_aba}' localmente: {err}")
        print(f"[ERRO SALVAR EXCEL]: {err}")

    sincronizar_google_sheets(df_limpo, nome_aba)
    return sucesso_local

def excluir_setor(setor_nome: str, unidade_nome: str) -> None:
    """Remove um setor inteiro (linha) da aba da unidade correspondente."""
    df, _ = carregar_dados_excel(unidade_nome)
    if df.empty or COLUNA_CHAVE not in df.columns:
        return
    mascara_manter = ~df[COLUNA_CHAVE].str.strip().str.lower().eq(setor_nome.strip().lower())
    df_filtrado = df[mascara_manter].copy()
    salvar_no_excel(df_filtrado, unidade_nome)
    st.session_state.df_historico = df_filtrado

def excluir_patrimonio(setor_nome: str, coluna_patrimonio: str, unidade_nome: str) -> None:
    """Limpa a célula do patrimônio e seu respectivo fabricante na aba correspondente."""
    df, _ = carregar_dados_excel(unidade_nome)
    if df.empty or COLUNA_CHAVE not in df.columns or coluna_patrimonio not in df.columns:
        return
    mascara_setor = df[COLUNA_CHAVE].str.strip().str.lower().eq(setor_nome.strip().lower())
    if mascara_setor.any():
        df.loc[mascara_setor, coluna_patrimonio] = ""
        col_fab = f"Fabricante {coluna_patrimonio}"
        if col_fab in df.columns:
            df.loc[mascara_setor, col_fab] = ""
        salvar_no_excel(df, unidade_nome)
        st.session_state.df_historico = df